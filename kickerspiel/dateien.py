"""Excel-Dateien je Spieltag: schreiben (zum Prüfen) und wieder lesen (als Eingabe
für den nächsten Schritt). Die Excel-Dateien sind die Schnittstelle zwischen den
Stufen – Korrekturen macht man dort, dann läuft die nächste Stufe neu.

  aufstellungen_stXX.xlsx  Blatt "Aufstellungen", "Mails"
  kickerdaten_stXX.xlsx    Blatt "Spieler", "Vereine", "Tore", "Prüfung"
  auswertung_stXX.xlsx     Blatt "Spieltag", "Elf je Manager", "Protokoll"
  saison.xlsx              Blatt "Tabelle", "Verlauf", "Kategorie-Profil", "Spieler-Beiträge"
"""
from __future__ import annotations

import json
from fractions import Fraction
from pathlib import Path
from typing import Iterable, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import KATEGORIEN, Herkunft, Kategorie, ManagerErgebnis, Saison, SpieltagErgebnis, punkte_text, schnitt_text
from .kicker_parser import KickerImport
from .mail_parser import MailAufstellung
from .model import Aufstellung, Spielerdaten, Spieltagsdaten, Vereinsdaten, ist_unbekannt, unbekannt_aufloesen
from .spielerbasis import Spielerbasis

GELB = PatternFill("solid", fgColor="FFF2CC")
ROT = PatternFill("solid", fgColor="F8CBAD")
GRUEN = PatternFill("solid", fgColor="E2EFDA")


def _blatt(wb, titel: str, kopf: list[str], breiten: Optional[list[int]] = None):
    ws = wb.create_sheet(titel)
    ws.append(kopf)
    for c in range(1, len(kopf) + 1):
        ws.cell(1, c).font = Font(bold=True)
        ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(breiten or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _fertig(wb, pfad: Path) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    for ws in wb.worksheets:
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)


def _f(x: Fraction) -> float:
    return float(x)


# ---------------------------------------------------------------------------
# Aufstellungen
# ---------------------------------------------------------------------------

def aufstellungen_schreiben(pfad: Path, mails: list[MailAufstellung], basis: Spielerbasis, spieltag: int) -> None:
    wb = openpyxl.Workbook()
    ws = _blatt(wb, "Aufstellungen", ["Manager", "Team", "Rolle", "Nr", "Name laut Mail", "Spieler-ID", "kicker-Name", "Verein", "Position", "Zuordnung", "Hinweis"],
                [11, 14, 7, 5, 18, 26, 18, 12, 9, 10, 60])
    for ma in mails:
        if not ma.aufstellung:
            continue
        team = basis.teams.get(ma.manager, "")
        nr_s = nr_b = 0
        zuo = {(z.rolle, i): z for i, z in enumerate(ma.zuordnungen)}
        stamm = [z for z in ma.zuordnungen if z.rolle == "Stamm"]
        bank = [z for z in ma.zuordnungen if z.rolle == "Bank"]
        for sid, z in zip(ma.aufstellung.start, stamm):
            nr_s += 1
            _zeile(ws, ma.manager, team, "Stamm", nr_s, z, sid, basis)
        for z in bank:
            if z.spieler is None:
                _zeile(ws, ma.manager, team, "Bank", None, z, "", basis)
        for sid, z in zip(ma.aufstellung.bank, [z for z in bank if z.spieler]):
            nr_b += 1
            _zeile(ws, ma.manager, team, "Bank", nr_b, z, sid, basis)
    ws2 = _blatt(wb, "Mails", ["Manager", "Datei", "Absender", "Datum", "Betreff", "Fehler", "Warnungen", "Ignorierter Text"], [11, 40, 28, 22, 34, 40, 50, 60])
    for ma in mails:
        ws2.append([ma.manager or "?", ma.datei, ma.absender, ma.datum, ma.betreff, "; ".join(ma.fehler), "; ".join(ma.warnungen), " | ".join(ma.ignoriert)])
        if ma.fehler:
            ws2.cell(ws2.max_row, 6).fill = ROT
    ws3 = _blatt(wb, "Hinweise", ["Text"], [140])
    for t in [
        f"Aufstellungen {spieltag}. Spieltag – erzeugt aus den .eml-Dateien. Dieses Blatt ist die Eingabe für 'auswerten'.",
        "Zuordnung: exakt = Name eindeutig im Kader; tolerant = Schreibweise korrigiert (gelb); unbekannt = kein Kaderspieler passt (rot) – Spieler-ID '?POS:Name' bleibt als Strafplatz stehen.",
        "Korrigieren: Spieler-ID ändern (Format 'Verein:kicker-Kurzname' wie in der Spielerbasis) oder Zeile löschen; Rolle Stamm/Bank und Nr bestimmen die Reihenfolge. Danach 'auswerten' neu laufen lassen.",
    ]:
        ws3.append([t])
    _fertig(wb, pfad)


def _zeile(ws, manager, team, rolle, nr, z, sid, basis: Spielerbasis) -> None:
    sp = z.spieler
    ws.append([manager, team, rolle, nr, z.name_mail, sid, sp.name if sp else "", sp.verein if sp else "", sp.position.value if sp else (unbekannt_aufloesen(sid)[0].value if sid and ist_unbekannt(sid) else ""), z.art, z.hinweis])
    if z.art == "tolerant":
        ws.cell(ws.max_row, 10).fill = GELB
    elif z.art == "unbekannt":
        for c in (5, 6, 10):
            ws.cell(ws.max_row, c).fill = ROT


def aufstellungen_lesen(pfad: Path, basis: Spielerbasis) -> list[Aufstellung]:
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb["Aufstellungen"]
    rows = list(ws.iter_rows(values_only=True))
    kopf = [str(c) for c in rows[0]]
    idx = {k: i for i, k in enumerate(kopf)}
    je_manager: dict[str, dict] = {}
    for r in rows[1:]:
        if not r or not r[idx["Manager"]]:
            continue
        m = str(r[idx["Manager"]]).strip()
        sid = str(r[idx["Spieler-ID"]] or "").strip()
        rolle = str(r[idx["Rolle"]] or "").strip()
        nr = r[idx["Nr"]]
        if not sid:
            continue
        d = je_manager.setdefault(m, {"start": [], "bank": []})
        try:
            nr = int(nr) if nr not in (None, "") else 10**6
        except (TypeError, ValueError):
            nr = 10**6
        d["start" if rolle == "Stamm" else "bank"].append((nr, sid))
    ergebnis = []
    for m, d in sorted(je_manager.items()):
        start = [s for _, s in sorted(d["start"], key=lambda t: t[0])]
        bank = [s for _, s in sorted(d["bank"], key=lambda t: t[0])]
        ergebnis.append(Aufstellung(m, start, bank))
    return ergebnis


# ---------------------------------------------------------------------------
# kicker-Daten
# ---------------------------------------------------------------------------

def kickerdaten_schreiben(pfad: Path, imp: KickerImport, basis: Spielerbasis) -> None:
    wb = openpyxl.Workbook()
    ws = _blatt(wb, "Spieler", ["Verein", "kicker-Name", "Status", "Minute", "Note", "Tore", "Vorlagen", "Elf des Tages", "Spieler-ID", "Manager", "Position"],
                [12, 24, 13, 8, 7, 6, 9, 13, 26, 11, 9])
    edt = set(imp.edt)
    for s in sorted(imp.alle_spieler(), key=lambda s: (s.verein, s.status != "Startelf", s.name)):
        sid = imp.zuordnung.get((s.verein, s.name), "")
        sp = basis.spieler.get(sid)
        ws.append([s.verein, s.name, s.status, s.minute or "", s.note or "", s.tore or None, s.vorlagen or None, "ja" if (s.verein, s.name) in edt else "", sid, sp.manager if sp else "", sp.position.value if sp else ""])
        if sid:
            ws.cell(ws.max_row, 9).fill = GRUEN
    ws2 = _blatt(wb, "Vereine", ["Verein", "Gegner", "Heim/Gast", "Ergebnis", "Gegentore", "Gespielt", "Quelle"], [12, 12, 9, 9, 10, 9, 40])
    for sp in imp.spiele:
        ws2.append([sp.heim, sp.gast, "Heim", f"{sp.tore_heim}:{sp.tore_gast}", sp.tore_gast, "ja", sp.quelle])
        ws2.append([sp.gast, sp.heim, "Gast", f"{sp.tore_heim}:{sp.tore_gast}", sp.tore_heim, "ja", sp.quelle])
    ws3 = _blatt(wb, "Tore", ["Spiel", "Minute", "Stand", "Tor für", "Schütze", "Eigentor", "Art", "Vorlage"], [26, 8, 7, 12, 20, 8, 16, 20])
    for sp in imp.spiele:
        for t in sp.tore:
            ws3.append([f"{sp.heim} – {sp.gast}", t.minute, t.stand, t.verein, t.schuetze, "ja" if t.eigentor else "", t.art, t.vorlage or ""])
    ws4 = _blatt(wb, "Elf des Tages", ["Verein", "kicker-Name", "Spieler-ID"], [12, 24, 26])
    for v, n in imp.edt:
        ws4.append([v, n, imp.zuordnung.get((v, n), "")])
    ws5 = _blatt(wb, "Prüfung", ["Quelle", "Hinweis"], [30, 100])
    for sp in imp.spiele:
        for w in sp.warnungen:
            ws5.append([sp.quelle, w])
    for w in imp.warnungen:
        ws5.append(["Zuordnung", w])
    fehlend = [sp for sid, sp in basis.spieler.items() if sid not in imp.zuordnung.values() and sp.verein in {v for s in imp.spiele for v in (s.heim, s.gast)}]
    for sp in fehlend:
        ws5.append(["Nicht im kicker-Kader", f"{sp.name} ({sp.verein}, {sp.manager}) taucht im Schema nicht auf → 0 Minuten"])
    ws6 = _blatt(wb, "Hinweise", ["Text"], [140])
    for t in [
        "kicker-Daten – erzeugt aus den gespeicherten Schema-Seiten (daten/spieltag_XX/kicker/). Dieses Blatt ist die Eingabe für 'auswerten'.",
        "Blatt Spieler: alle Spieler aller Spiele; Spieler-ID (grün) = in der Spielerbasis gefunden. Note leer = keine Note. Status Reservebank = 0 Minuten (Strafgegentor bei Strafnote).",
        "Korrigieren: Note/Tore/Vorlagen/Elf des Tages direkt ändern; fehlende Zuordnung durch Eintragen der Spieler-ID ergänzen. Blatt Vereine: Gegentore ggf. anpassen. Danach 'auswerten' neu laufen lassen.",
    ]:
        ws6.append([t])
    _fertig(wb, pfad)


def kickerdaten_lesen(pfad: Path, spieltag: int, basis: Spielerbasis) -> Spieltagsdaten:
    wb = openpyxl.load_workbook(pfad, data_only=True)
    daten = Spieltagsdaten(spieltag)
    rows = list(wb["Vereine"].iter_rows(values_only=True))
    idx = {str(k): i for i, k in enumerate(rows[0])}
    for r in rows[1:]:
        if r and r[idx["Verein"]]:
            daten.vereine[str(r[idx["Verein"]])] = Vereinsdaten(int(r[idx["Gegentore"]] or 0), str(r[idx["Gespielt"]] or "ja").lower() != "nein")
    rows = list(wb["Spieler"].iter_rows(values_only=True))
    idx = {str(k): i for i, k in enumerate(rows[0])}
    for r in rows[1:]:
        if not r or not r[idx["Spieler-ID"]]:
            continue
        sid = str(r[idx["Spieler-ID"]]).strip()
        status = str(r[idx["Status"]] or "")
        note = r[idx["Note"]]
        edt = str(r[idx["Elf des Tages"]] or "").strip().lower() in ("ja", "x", "1", "true")
        daten.spieler[sid] = Spielerdaten(note if note not in (None, "") else None, int(r[idx["Tore"]] or 0), int(r[idx["Vorlagen"]] or 0), edt,
                                          eingesetzt=(status != "Reservebank"))
    for sid, sp in basis.spieler.items():
        if sid not in daten.spieler and sp.verein in daten.vereine:
            daten.spieler[sid] = Spielerdaten(None, 0, 0, False, eingesetzt=False)
    return daten


# ---------------------------------------------------------------------------
# Auswertung und Saison
# ---------------------------------------------------------------------------

def auswertung_schreiben(pfad: Path, erg: SpieltagErgebnis, basis: Spielerbasis) -> None:
    wb = openpyxl.Workbook()
    ws = _blatt(wb, "Spieltag", ["Platz", "Manager", "Team", "Summe", "Ø-Note", "RP Note", "Gegentore", "RP Gegentore", "Tore", "RP Tore", "Vorlagen", "RP Vorlagen", "Elf des Tages", "RP EdT", "Spieltagssieg"],
                [6, 11, 14, 8, 8, 8, 10, 12, 6, 8, 9, 11, 12, 8, 12])
    for m in erg.reihenfolge:
        e = erg.manager[m]
        ws.append([e.platz, m, basis.teams.get(m, ""), _f(e.summe), round(_f(e.werte[Kategorie.NOTE]), 2), _f(e.rangpunkte[Kategorie.NOTE]),
                   int(e.werte[Kategorie.GEGENTORE]), _f(e.rangpunkte[Kategorie.GEGENTORE]), int(e.werte[Kategorie.TORE]), _f(e.rangpunkte[Kategorie.TORE]),
                   int(e.werte[Kategorie.VORLAGEN]), _f(e.rangpunkte[Kategorie.VORLAGEN]), int(e.werte[Kategorie.EDT]), _f(e.rangpunkte[Kategorie.EDT]), _f(erg.spieltagssieg[m])])
    ws2 = _blatt(wb, "Elf je Manager", ["Manager", "Pos", "Spieler", "Verein", "Herkunft", "Ersetzt", "Bank", "Note", "Strafnote", "Eingesetzt", "Gegentore Verein", "Strafgegentor", "Gegentore gew.", "Tore", "Tore gew.", "Vorlagen", "Vorlagen gew.", "Elf des Tages"],
                 [11, 5, 20, 12, 11, 16, 5, 6, 9, 10, 10, 10, 10, 6, 8, 8, 10, 8])
    for m in erg.reihenfolge:
        for p in erg.manager[m].elf:
            ws2.append([m, p.position.value, p.name, p.spieler.verein if p.spieler else "", p.herkunft.value, p.ersetzt.name if p.ersetzt else "", p.bankplatz or "",
                        _f(p.note), "ja" if p.strafnote else "", "" if p.eingesetzt is None else ("ja" if p.eingesetzt else "nein"),
                        p.gegentore_verein if p.position.value in ("TOR", "ABW") else "", p.strafgegentore or "", p.gegentore_gew if p.position.value in ("TOR", "ABW") else "",
                        p.tore or "", p.tore_gew or "", p.vorlagen or "", p.vorlagen_gew or "", "ja" if p.edt else ""])
            if p.strafnote:
                ws2.cell(ws2.max_row, 8).fill = ROT
            if p.herkunft == Herkunft.NACHRUECKER:
                ws2.cell(ws2.max_row, 5).fill = GELB
    ws3 = _blatt(wb, "Protokoll", ["Manager", "Art", "Text"], [11, 10, 120])
    for m in erg.reihenfolge:
        for z in erg.manager[m].protokoll:
            ws3.append([m, "Protokoll", z.strip()])
        for w in erg.manager[m].warnungen:
            ws3.append([m, "Warnung", w])
            ws3.cell(ws3.max_row, 2).fill = GELB
    _fertig(wb, pfad)


def saison_schreiben(pfad: Path, saison: Saison, basis: Spielerbasis) -> None:
    wb = openpyxl.Workbook()
    ws = _blatt(wb, "Tabelle", ["Platz", "Manager", "Team", "Punkte", "Spieltagssiege", "Spieltage"], [6, 11, 14, 8, 13, 9])
    for z in saison.zeilen:
        ws.append([z.platz, z.manager, basis.teams.get(z.manager, ""), _f(z.punkte), _f(z.spieltagssiege), len(z.verlauf)])
    ws2 = _blatt(wb, "Verlauf", ["Manager"] + [f"ST {s}" for s in saison.spieltage] + ["Summe"] + [f"kum. ST {s}" for s in saison.spieltage], [11] + [7] * (2 * len(saison.spieltage) + 1))
    for z in saison.zeilen:
        ws2.append([z.manager] + [_f(z.verlauf.get(s, Fraction(0))) for s in saison.spieltage] + [_f(z.punkte)] + [_f(z.kumuliert.get(s, Fraction(0))) for s in saison.spieltage])
    ws3 = _blatt(wb, "Kategorie-Profil", ["Manager"] + [k.value for k in KATEGORIEN] + ["Summe"], [11, 8, 10, 8, 10, 13, 8])
    for z in saison.zeilen:
        ws3.append([z.manager] + [_f(z.kategorie_profil[k]) for k in KATEGORIEN] + [_f(z.punkte)])
    ws4 = _blatt(wb, "Spieler-Beiträge", ["Manager", "Spieler", "Verein", "Position", "Einsätze gewertet", "davon nachgerückt", "Strafnoten", "Ø-Note", "Tore", "Tore gew.", "Vorlagen", "Vorlagen gew.", "Gegentore gew.", "Elf des Tages"],
                 [11, 20, 12, 8, 10, 10, 9, 7, 6, 8, 8, 10, 10, 8])
    for b in sorted(saison.beitraege.values(), key=lambda b: (b.spieler.manager, b.spieler.position.value, b.spieler.name)):
        ws4.append([b.spieler.manager, b.spieler.name, b.spieler.verein, b.spieler.position.value, b.einsaetze, b.nachgerueckt, b.strafnoten,
                    round(_f(b.notenschnitt), 2) if b.notenschnitt is not None else "", b.tore, b.tore_gew, b.vorlagen, b.vorlagen_gew, b.gegentore_gew, b.edt])
    _fertig(wb, pfad)


# ---------------------------------------------------------------------------
# JSON: Ergebnis je Spieltag (für die Saisontabelle) und Rohdaten
# ---------------------------------------------------------------------------

def ergebnis_json(erg: SpieltagErgebnis) -> dict:
    return {
        "spieltag": erg.spieltag,
        "manager": {
            m: {
                "summe": str(e.summe), "platz": e.platz,
                "werte": {k.value: str(e.werte[k]) for k in KATEGORIEN},
                "rangpunkte": {k.value: str(e.rangpunkte[k]) for k in KATEGORIEN},
                "spieltagssieg": str(erg.spieltagssieg[m]),
                "elf": [{"position": p.position.value, "spieler_id": p.spieler.id if p.spieler else None, "name": p.name, "herkunft": p.herkunft.value,
                         "note": str(p.note), "strafnote": p.strafnote, "tore": p.tore, "tore_gew": p.tore_gew, "vorlagen": p.vorlagen, "vorlagen_gew": p.vorlagen_gew,
                         "gegentore_gew": p.gegentore_gew, "edt": p.edt} for p in e.elf],
                "protokoll": e.protokoll, "warnungen": e.warnungen,
            } for m, e in erg.manager.items()
        },
    }


def json_schreiben(pfad: Path, daten: dict) -> None:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    pfad.write_text(json.dumps(daten, ensure_ascii=False, indent=2), encoding="utf-8")
