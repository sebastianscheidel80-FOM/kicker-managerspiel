"""Spielerbasis einlesen (Excel oder CSV) und Namen zuordnen.

Die Spielerbasis ist die Wahrheit über Kader, Vereine und kicker-Positionen.
Spieler-ID = "<Verein>:<kicker-Kurzname>", z. B. "Frankfurt:R. Koch".
"""
from __future__ import annotations

import csv
import difflib
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

from .model import Position, Spieler

SPALTEN = {
    "manager": "Manager",
    "team": "Team (fiktiv)",
    "auktion": "Name laut Auktion",
    "kaderliste": "kicker-Name (Kaderliste)",
    "kurz": "kicker-Kurzname (Schema)",
    "verein": "Verein",
    "position": "Position",
    "preis": "Kaufpreis (Mio.)",
    "ab": "Gültig ab Spieltag",
    "bis": "Gültig bis Spieltag",
}


INITIAL = re.compile(r"^(?:[A-Za-zÄÖÜ][a-z]{0,2}\.(?:-[A-Za-z]\.)?\s*)+")


def ohne_initial(name: str) -> str:
    """'R. Koch' → 'Koch', 'J.-S. Lee' → 'Lee', 'Kai. Sano' → 'Sano', 'La. Günther' → 'Günther'."""
    return INITIAL.sub("", name.strip())


def normalisieren(text: str) -> str:
    """Kleinbuchstaben, ohne Akzente, ohne Satzzeichen – für den Namensvergleich."""
    t = unicodedata.normalize("NFKD", str(text))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = t.lower().replace("ß", "ss")
    return re.sub(r"[^a-z0-9 ]", " ", t).strip()


@dataclass
class Spielerbasis:
    spieler: dict[str, Spieler] = field(default_factory=dict)        # ID -> Spieler
    teams: dict[str, str] = field(default_factory=dict)              # Manager -> Team (fiktiv)
    aliase: dict[str, set[str]] = field(default_factory=dict)        # ID -> normalisierte Namensvarianten

    def manager(self) -> list[str]:
        return sorted({s.manager for s in self.spieler.values()})

    def kader(self, manager: str, spieltag: Optional[int] = None) -> list[Spieler]:
        return [s for s in self.spieler.values() if s.gehoert_zu(manager, spieltag)]

    def vereine(self) -> set[str]:
        return {s.verein for s in self.spieler.values()}

    # -- Namenszuordnung -----------------------------------------------------

    def finde(self, name: str, kandidaten: Iterable[Spieler], position: Optional[Position] = None) -> tuple[Optional[Spieler], str, list[Spieler]]:
        """Ordnet einen Namen einem Spieler aus `kandidaten` zu.

        Liefert (spieler, art, alternativen). art: "exakt", "tolerant" oder "unbekannt".
        Ein Initial vor dem Namen ("R. Koch") wird beim Vergleich toleriert.
        """
        n = normalisieren(name)
        n_ohne_initial = normalisieren(ohne_initial(name))
        exakt, tolerant = [], []
        for s in kandidaten:
            al = self.aliase.get(s.id, set())
            if n in al or n_ohne_initial in al:
                exakt.append(s)
                continue
            beste = max((difflib.SequenceMatcher(None, n_ohne_initial, a).ratio() for a in al), default=0)
            if beste >= 0.8:
                tolerant.append((beste, s))
        if position is not None:
            exakt_pos = [s for s in exakt if s.position == position]
            if exakt_pos:
                exakt = exakt_pos
        if len(exakt) == 1:
            return exakt[0], "exakt", []
        if len(exakt) > 1:
            return None, "unbekannt", exakt
        if tolerant:
            tolerant.sort(key=lambda t: -t[0])
            if len(tolerant) == 1 or tolerant[0][0] - tolerant[1][0] > 0.05:
                return tolerant[0][1], "tolerant", [t[1] for t in tolerant[1:3]]
            return None, "unbekannt", [t[1] for t in tolerant[:3]]
        return None, "unbekannt", []


def _aliase_fuer(kurz: str, kaderliste: str, auktion: str) -> set[str]:
    al = set()
    for variante in (kurz, kaderliste, auktion):
        if not variante:
            continue
        v = normalisieren(variante)
        al.add(v)
        al.add(normalisieren(ohne_initial(variante)))
    # Kaderlistenform "Nachname Vorname": Nachname allein, Vorname allein, Teil-Nachnamen
    toks = normalisieren(kaderliste).split()
    if len(toks) >= 2:
        al.add(toks[0])
        al.add(toks[-1])
        for k in range(1, len(toks)):
            al.add(" ".join(toks[:k]))
    al.discard("")
    return al


def lade_spielerbasis(pfad: str | Path) -> Spielerbasis:
    pfad = Path(pfad)
    zeilen = _zeilen_lesen(pfad)
    basis = Spielerbasis()
    for z in zeilen:
        manager = str(z.get(SPALTEN["manager"], "")).strip()
        kurz = str(z.get(SPALTEN["kurz"], "") or "").strip()
        kaderliste = str(z.get(SPALTEN["kaderliste"], "") or "").strip()
        auktion = str(z.get(SPALTEN["auktion"], "") or "").strip()
        verein = str(z.get(SPALTEN["verein"], "") or "").strip()
        if not manager or not (kurz or kaderliste):
            continue
        name = kurz or kaderliste.split()[0]
        position = Position.parse(z.get(SPALTEN["position"], ""))
        preis = z.get(SPALTEN["preis"])
        try:
            preis = float(str(preis).replace(",", ".")) if preis not in (None, "") else None
        except ValueError:
            preis = None
        ab = int(float(z.get(SPALTEN["ab"]) or 1))
        bis_raw = z.get(SPALTEN["bis"])
        bis = int(float(bis_raw)) if bis_raw not in (None, "") else None
        sid = f"{verein}:{name}"
        if sid in basis.spieler:
            raise ValueError(f"Doppelte Spieler-ID in der Spielerbasis: {sid}")
        basis.spieler[sid] = Spieler(sid, name, verein, position, manager, preis, ab, bis)
        basis.aliase[sid] = _aliase_fuer(kurz, kaderliste, auktion)
        team = str(z.get(SPALTEN["team"], "") or "").strip()
        if team:
            basis.teams[manager] = team
    return basis


def _zeilen_lesen(pfad: Path) -> list[dict]:
    if pfad.suffix.lower() == ".csv":
        with open(pfad, encoding="utf-8", newline="") as f:
            return list(csv.DictReader(f, delimiter=";"))
    import openpyxl

    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb["Spielerbasis"] if "Spielerbasis" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    kopf = [str(c).strip() if c is not None else "" for c in rows[0]]
    return [dict(zip(kopf, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
